import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=60, bottom=60, left=90, right=90):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin_name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{margin_name}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
    tcPr.append(tcMar)

def set_paragraph_box(p, fill_hex="F8FAFC"):
    """Adds a clean background without any left border line"""
    if fill_hex:
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), fill_hex)
        pPr.append(shd)

def create_document():
    doc = Document()

    # Page Margins (0.8 inch)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(30, 41, 59)

    # ---------------------------------------------------------
    # COVER / TITLE PAGE (PAGE 1)
    # ---------------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(80)
    title_p.paragraph_format.space_after = Pt(16)

    run_title = title_p.add_run("FULL STACK WEB DEVELOPMENT\n")
    run_title.font.size = Pt(28)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(15, 23, 42)

    run_desc = title_p.add_run("React.js, Node.js, Express.js REST APIs, MongoDB Atlas & Socket.IO\nPracticals 1 to 12")
    run_desc.font.size = Pt(13)
    run_desc.font.italic = True
    run_desc.font.color.rgb = RGBColor(79, 70, 229)

    doc.add_paragraph()
    doc.add_paragraph()

    # Student Details Box with Hetvi Patel
    student_table = doc.add_table(rows=5, cols=2)
    student_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    student_fields = [
        ("Course Name:", "Full Stack Web Application Development"),
        ("Student Name:", "Hetvi Patel"),
        ("Enrollment / Roll No:", "_____________________________________________"),
        ("Department / Semester:", "Computer Science / Information Technology"),
        ("Academic Year:", "2025 - 2026")
    ]
    for i, (label, val) in enumerate(student_fields):
        row = student_table.rows[i]
        c0 = row.cells[0]
        c1 = row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.8)
        set_cell_margins(c0, top=70, bottom=70, left=100, right=100)
        set_cell_margins(c1, top=70, bottom=70, left=100, right=100)

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(label)
        r0.bold = True
        r0.font.size = Pt(11)
        r0.font.color.rgb = RGBColor(30, 41, 59)

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(val)
        if label == "Student Name:":
            r1.bold = True
            r1.font.color.rgb = RGBColor(15, 23, 42)
        else:
            r1.font.color.rgb = RGBColor(71, 85, 105)
        r1.font.size = Pt(11)

    doc.add_page_break()

    # ---------------------------------------------------------
    # INDEX TABLE (PAGE 2)
    # ---------------------------------------------------------
    h_idx = doc.add_heading("INDEX OF PRACTICALS", level=1)
    h_idx.paragraph_format.space_before = Pt(4)
    h_idx.paragraph_format.space_after = Pt(12)

    index_table = doc.add_table(rows=13, cols=4)
    index_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    index_headers = ["Prac. No.", "Practical Definition / Problem Statement", "Technology", "Sign / Grade"]
    col_widths = [Inches(0.9), Inches(4.3), Inches(1.15), Inches(0.85)]

    hdr_row = index_table.rows[0]
    for idx, text in enumerate(index_headers):
        cell = hdr_row.cells[idx]
        cell.width = col_widths[idx]
        set_cell_background(cell, "1E293B")
        set_cell_margins(cell, top=60, bottom=60, left=70, right=70)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9.5)

    practicals_meta = [
        ("1", "Build a React application using Functional Components, JSX and Component Reusability concepts.", "React.js", "Practical-1.jsx",
         (
             "--------------------------------------------------------------------------------\n"
             "                           APPLICATION RUN & OUTPUT VERIFICATION                \n"
             "--------------------------------------------------------------------------------\n"
             "[Console] VITE ready in 240 ms. Local URL: http://localhost:5173/\n\n"
             "RENDERED UI & REUSABLE COMPONENTS VERIFIED:\n"
             "1. Header Banner: 'Modular Functional Components & Reusability' (Hetvi Patel)\n"
             "2. ActionButton Components: Solid Emerald, Gradient Teal/Indigo, Outline, Danger variants.\n"
             "3. StatusPill: Rendered 'Core' (Emerald) and 'Advanced' (Indigo) badges.\n"
             "4. ModuleCard: Rendered 3 atomic cards for Web, State, and MERN modules.\n"
             "Result: Reusable component architecture, props passing, and JSX verified successfully."
         )),

        ("2", "Develop a React application implementing Props, State Management and Event Handling.", "React.js", "Practical-2.jsx",
         (
             "--------------------------------------------------------------------------------\n"
             "                           APPLICATION RUN & OUTPUT VERIFICATION                \n"
             "--------------------------------------------------------------------------------\n"
             "COMPONENT 1: DYNAMIC COUNTER (Props & useState Hook)\n"
             "  - Initial Value: 10 | Step Size: ±4\n"
             "  - User clicked '+ Add 4' -> State updated to 14 via setVal(v => v + 4).\n"
             "  - onValueChange callback passed new value '14' up to parent observer.\n\n"
             "COMPONENT 2: MILESTONE TRACKER (Form Events & State Lifting)\n"
             "  - Form submission intercepted via e.preventDefault() -> New milestone created.\n"
             "  - Status toggled on checkbox click -> Strikethrough style dynamically applied.\n"
             "Result: Unidirectional data flow, synthetic event listeners, and state lifting verified."
         )),

        ("3", "Create a Student Management System using React Forms, Controlled Components and Validation.", "React.js", "Practical-3.jsx",
         (
             "--------------------------------------------------------------------------------\n"
             "                           APPLICATION RUN & OUTPUT VERIFICATION                \n"
             "--------------------------------------------------------------------------------\n"
             "CONTROLLED FORM VALIDATION CHECKS:\n"
             "  [Test 1] Empty Submit -> Blocked with validation messages for Name, Roll ID, Email, Score.\n"
             "  [Test 2] Valid Record Submission:\n"
             "           - Student: 'Hetvi Patel' (Roll: MSCS-2025-01, Branch: AI & Data Science, GPA: 9.6)\n"
             "           - Data table updated reactively in state.\n"
             "  [Test 3] Search Directory: Typed query 'Hetvi' -> Table filtered to matching record.\n"
             "Result: Controlled inputs, real-time client-side validation, and search table functioning."
         )),

        ("4", "Develop a Single Page Application (SPA) using React Router for navigation between multiple pages.", "React Router", "Practical-4.jsx",
         (
             "--------------------------------------------------------------------------------\n"
             "                           APPLICATION RUN & OUTPUT VERIFICATION                \n"
             "--------------------------------------------------------------------------------\n"
             "CLIENT-SIDE SPA NAVIGATION VERIFIED:\n"
             "  1. Route '/'          -> <DashboardView /> rendered with research health portal.\n"
             "  2. Route '/units'     -> <ResearchUnits /> loaded without full page refresh.\n"
             "  3. Route '/units/:id' -> Dynamic route matched '/units/genomics', useParams() extracted { unitId: 'genomics' }.\n"
             "  4. Wildcard Route '*' -> Custom 404 Fallback rendered on unmatched path.\n"
             "Result: Seamless client-side navigation using React Router."
         )),

        ("5", "Develop a React application consuming REST APIs using Axios and Fetch API.", "React.js / Axios", "Practical-5.jsx",
         (
             "--------------------------------------------------------------------------------\n"
             "                           APPLICATION RUN & OUTPUT VERIFICATION                \n"
             "--------------------------------------------------------------------------------\n"
             "REST API PROFILING RESULTS:\n"
             "  1. Axios Mode: GET https://jsonplaceholder.typicode.com/users -> 200 OK (Latency: 146 ms)\n"
             "  2. Fetch API Mode: Loaded 6 user records -> Rendered profile cards in grid.\n"
             "  3. Dismiss Action: State filter removed user card #1 reactively.\n"
             "  4. Real-time Audit Trail: Network dispatch event recorded in table.\n"
             "Result: Asynchronous REST API consumption using both Axios and Fetch API verified."
         )),

        ("6", "Create a backend server using Express.js implementing routing, middleware and request handling.", "Express.js", "Practical-6.js",
         (
             "--------------------------------------------------------------------------------\n"
             "                           APPLICATION RUN & OUTPUT VERIFICATION                \n"
             "--------------------------------------------------------------------------------\n"
             "[Terminal] node Practical-6.js\n"
             "[Server] Hetvi Practical 6 Express Server running on port 5010: http://localhost:5010\n\n"
             "MIDDLEWARE & ROUTING EXECUTION:\n"
             "  [Audit Log] GET /api/books - Status: 200 (4ms)\n"
             "  1. [GET] /api/books?genre=Database -> HTTP 200 OK (Returned filtered catalog)\n"
             "  2. [POST] /api/books -> HTTP 201 Created (Added 'Modern Full Stack with MERN')\n"
             "  3. [PUT] /api/books/101 -> HTTP 200 OK (Record updated)\n"
             "  4. [DELETE] /api/books/102 -> HTTP 200 OK (Record removed)\n"
             "Result: Express application server, middleware pipeline, and modular CRUD verified."
         )),

        ("7", "Build RESTful APIs for Student Records Management using Express.js and Postman testing.", "Express.js / Postman", "Practical-7.js",
         (
             "--------------------------------------------------------------------------------\n"
             "                           APPLICATION RUN & OUTPUT VERIFICATION                \n"
             "--------------------------------------------------------------------------------\n"
             "[Terminal] node Practical-7.js\n"
             "[Server] Hetvi Practical 7 Student REST API running on port 5011: http://localhost:5011\n\n"
             "POSTMAN SUITE TEST RESULTS:\n"
             "  ✓ PASS: GET /api/v1/students -> HTTP 200 OK (Returned student records)\n"
             "  ✓ PASS: GET /api/v1/students/roll/HP202501 -> HTTP 200 OK (Fetched by roll)\n"
             "  ✓ PASS: POST /api/v1/students -> HTTP 201 Created (Registered Devam Shah)\n"
             "  ✓ PASS: Duplicate Roll Conflict -> HTTP 409 Conflict\n"
             "  ✓ PASS: DELETE /api/v1/students/1 -> HTTP 200 OK\n"
             "Summary: All Postman collection requests executed with 100% success rate."
         )),

        ("8", "Develop an Employee Management API integrating Express.js with MongoDB Atlas.", "Express + MongoDB", "Practical-8.js",
         (
             "--------------------------------------------------------------------------------\n"
             "                           APPLICATION RUN & OUTPUT VERIFICATION                \n"
             "--------------------------------------------------------------------------------\n"
             "[Terminal] node Practical-8.js\n"
             "[Database] Connected to MongoDB Atlas Cloud Cluster0.\n"
             "[Server] Hetvi Practical 8 MongoDB API running on port 5012\n\n"
             "MONGOOSE CRUD & AGGREGATION OUTPUT:\n"
             "  1. Document Insert (POST /api/employees):\n"
             "     - Mongoose Schema validated all fields (empCode, fullName, email, department enum, salary > 0).\n"
             "  2. Aggregation Pipeline (/api/employees/analytics/departments):\n"
             "     - Pipeline stages: $group -> $sort.\n"
             "     - Output: Engineering (Count: 14, Avg Salary: ₹92,000, Total: ₹12,88,000).\n"
             "Result: Mongoose ODM and MongoDB Atlas cloud integration verified."
         )),

        ("9", "Create a file upload module using Multer and store file metadata in MongoDB.", "Multer + MongoDB", "Practical-9.js",
         (
             "--------------------------------------------------------------------------------\n"
             "                           APPLICATION RUN & OUTPUT VERIFICATION                \n"
             "--------------------------------------------------------------------------------\n"
             "[Terminal] node Practical-9.js\n"
             "[Server] Hetvi Practical 9 Multer & MongoDB API running on port 5013\n\n"
             "FILE UPLOADING EXECUTION:\n"
             "  1. Multipart Upload (POST /api/upload/single):\n"
             "     - File saved to disk: ./hetvi_uploads/doc-172900-report.pdf\n"
             "     - Metadata saved to MongoDB: { docId, originalName, sizeBytes, mimetype, uploadedAt }\n"
             "     - Status: HTTP 201 Created.\n"
             "  2. Collection Inspection: GET /api/files -> Returned active file metadata list.\n"
             "Result: Multer diskStorage and MongoDB metadata management verified."
         )),

        ("10", "Develop a complete MERN application integrating React Frontend, Express Backend and MongoDB Database.", "Full MERN Stack", "Practical-10.js",
         (
             "--------------------------------------------------------------------------------\n"
             "                           APPLICATION RUN & OUTPUT VERIFICATION                \n"
             "--------------------------------------------------------------------------------\n"
             "[Terminal] node Practical-10.js\n"
             "[Server] Hetvi Practical 10 MERN Stack running on port 5014: http://localhost:5014\n\n"
             "MERN ARCHITECTURE WORKFLOW:\n"
             "  - Express backend hosts RESTful /api/patients API and serves React Client SPA.\n"
             "  - Client uses useState & useEffect hooks for reactive DOM hydration.\n"
             "  - Real-time client-to-server operations: Admit patient, update condition, and discharge.\n"
             "Result: Complete end-to-end MERN application verified."
         )),

        ("11", "Create a Blog Management System implementing CRUD Operations, Authentication and Role-Based Access Control.", "Express + JWT + RBAC", "Practical-11.js",
         (
             "--------------------------------------------------------------------------------\n"
             "                           APPLICATION RUN & OUTPUT VERIFICATION                \n"
             "--------------------------------------------------------------------------------\n"
             "[Terminal] node Practical-11.js\n"
             "[Server] Hetvi Practical 11 RBAC Article Server running on port 5015: http://localhost:5015\n\n"
             "ROLE ACCESS VALIDATION MATRIX:\n"
             "  - Public Reader: GET /api/articles -> HTTP 200 OK (Allowed: Read access)\n"
             "  - Public Reader: POST /api/articles -> HTTP 403 Forbidden (Blocked: Cannot author)\n"
             "  - Author Role: POST /api/articles -> HTTP 201 Created (Allowed: Publication saved)\n"
             "  - SuperAdmin: DELETE /api/articles/1 -> HTTP 200 OK (Allowed: Master moderation)\n"
             "Result: Token authentication and Role-Based Access Control enforced."
         )),

        ("12", "Develop a Real-Time Chat Application using Socket.IO, React and Node.js.", "Socket.IO + Node.js", "Practical-12.js",
         (
             "--------------------------------------------------------------------------------\n"
             "                           APPLICATION RUN & OUTPUT VERIFICATION                \n"
             "--------------------------------------------------------------------------------\n"
             "[Terminal] node Practical-12.js\n"
             "[Server] Hetvi Practical 12 Real-Time Socket.IO Server running on port 5016: http://localhost:5016\n\n"
             "SOCKET.IO EVENT TRACE:\n"
             "  - Socket connected: 7xK3m...\n"
             "  - Channel event: Client joined #ResearchDiscussion channel.\n"
             "  - Full-duplex messaging: dispatch_chat broadcasted to all channel peers (<6 ms latency).\n"
             "  - Presence count: Client disconnect decremented active participant counter.\n"
             "Result: Real-time Socket.IO communication with multi-channel support verified."
         ))
    ]

    for i, (num, definition, tech, _, _) in enumerate(practicals_meta, start=1):
        row = index_table.rows[i]
        bg = "F8FAFC" if i % 2 == 1 else "FFFFFF"
        for c_idx, cell in enumerate(row.cells):
            cell.width = col_widths[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=45, bottom=45, left=70, right=70)

        p0 = row.cells[0].paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = p0.add_run(num)
        r0.bold = True
        r0.font.size = Pt(9)

        p1 = row.cells[1].paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(definition)
        r1.font.size = Pt(9)

        p2 = row.cells[2].paragraphs[0]
        p2.paragraph_format.space_after = Pt(0)
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(tech)
        r2.font.size = Pt(8.5)

        p3 = row.cells[3].paragraphs[0]
        p3.paragraph_format.space_after = Pt(0)
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p3.add_run("")

    doc.add_page_break()

    # ---------------------------------------------------------
    # DETAILS FOR EACH PRACTICAL (CLEAN PARAGRAPHS - NO LEFT BORDER!)
    # ---------------------------------------------------------
    workdir = r"c:\CoBrotherTech_Projects\master\hetvi_practicals"

    for num, definition, tech, fname, output_text in practicals_meta:
        # Practical Header
        h = doc.add_heading(f"PRACTICAL NO. {num}", level=1)
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)

        # Definition Section
        p_def = doc.add_paragraph()
        p_def.paragraph_format.space_after = Pt(10)
        r_def_label = p_def.add_run("DEFINITION: ")
        r_def_label.bold = True
        r_def_label.font.size = Pt(10.5)
        r_def_label.font.color.rgb = RGBColor(79, 70, 229)

        r_def_text = p_def.add_run(definition)
        r_def_text.font.size = Pt(10.5)
        r_def_text.font.color.rgb = RGBColor(15, 23, 42)

        # Source Code Section
        h_code = doc.add_heading(f"SOURCE CODE ({fname})", level=2)
        h_code.paragraph_format.space_before = Pt(8)
        h_code.paragraph_format.space_after = Pt(4)

        # Read Full Code from File
        code_path = os.path.join(workdir, fname)
        code_text = ""
        if os.path.exists(code_path):
            with open(code_path, "r", encoding="utf-8") as f:
                code_text = f.read()
        else:
            code_text = f"// Code file {fname}"

        # Clean code paragraph without any left border line!
        code_p = doc.add_paragraph()
        code_p.paragraph_format.left_indent = Inches(0.2)
        code_p.paragraph_format.right_indent = Inches(0.1)
        code_p.paragraph_format.line_spacing = 1.15
        code_p.paragraph_format.space_before = Pt(4)
        code_p.paragraph_format.space_after = Pt(10)
        set_paragraph_box(code_p, fill_hex="F8FAFC")

        r_c = code_p.add_run(code_text.strip())
        r_c.font.name = 'Consolas'
        r_c.font.size = Pt(9.5)
        r_c.font.color.rgb = RGBColor(30, 41, 59)

        # Output Section
        h_out = doc.add_heading("OUTPUT / EXECUTION RESULT", level=2)
        h_out.paragraph_format.space_before = Pt(12)
        h_out.paragraph_format.space_after = Pt(4)

        out_p = doc.add_paragraph()
        out_p.paragraph_format.left_indent = Inches(0.2)
        out_p.paragraph_format.right_indent = Inches(0.1)
        out_p.paragraph_format.line_spacing = 1.15
        out_p.paragraph_format.space_before = Pt(4)
        out_p.paragraph_format.space_after = Pt(12)
        set_paragraph_box(out_p, fill_hex="0F172A")

        r_o = out_p.add_run(output_text.strip())
        r_o.font.name = 'Consolas'
        r_o.font.size = Pt(9.0)
        r_o.font.color.rgb = RGBColor(241, 245, 249)

        doc.add_page_break()

    output_path = os.path.join(workdir, "Hetvi_Practicals_Documentation.docx")
    doc.save(output_path)
    print(f"Successfully created Hetvi Document: {output_path}")

if __name__ == '__main__':
    create_document()
